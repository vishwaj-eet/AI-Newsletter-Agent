
from io import BytesIO
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Inches, Pt
from xhtml2pdf import pisa


def get_newsletter_html(result):
    """Get newsletter HTML from the agent result or saved HTML file."""
    for key in ("newsletter_html", "html", "newsletter"):
        value = result.get(key)
        if isinstance(value, str) and value.strip():
            return value

    output_path = result.get("output_path")
    if output_path and Path(output_path).exists():
        return Path(output_path).read_text(encoding="utf-8")

    raise ValueError("Could not find newsletter HTML in the agent result.")



from io import BytesIO
from xhtml2pdf import pisa

def create_pdf(html_content: str) -> bytes:
    """Convert newsletter HTML into a properly paginated A4 PDF."""

    pdf_buffer = BytesIO()

    styled_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            @page {{
                size: A4;
                margin: 20mm 18mm 20mm 18mm;
            }}

            body {{
                font-family: Helvetica, Arial, sans-serif;
                font-size: 10pt;
                line-height: 1.5;
                color: #222222;
                margin: 0;
                padding: 0;
                width: 100%;
            }}

            h1 {{
                font-size: 22pt;
                margin-bottom: 12pt;
                line-height: 1.2;
            }}

            h2 {{
                font-size: 15pt;
                margin-top: 18pt;
                margin-bottom: 8pt;
                line-height: 1.3;
            }}

            h3 {{
                font-size: 12pt;
                margin-top: 12pt;
            }}

            p {{
                margin: 0 0 10pt 0;
                overflow-wrap: break-word;
            }}

            img {{
                max-width: 100%;
                height: auto;
            }}

            a {{
                color: #175cd3;
                word-wrap: break-word;
            }}

            table {{
                width: 100%;
                border-collapse: collapse;
            }}

            th, td {{
                padding: 6pt;
                border: 1px solid #dddddd;
                text-align: left;
            }}

            ul, ol {{
                margin-bottom: 10pt;
                padding-left: 20pt;
            }}

            .article, section {{
                margin-bottom: 18pt;
                page-break-inside: auto;
            }}

            h1, h2, h3 {{
                page-break-after: avoid;
            }}
        </style>
    </head>
    <body>
        {html_content}
    </body>
    </html>
    """

    result = pisa.CreatePDF(
        src=styled_html,
        dest=pdf_buffer,
        encoding="utf-8",
    )

    if result.err:
        raise RuntimeError("Failed to generate the newsletter PDF.")

    return pdf_buffer.getvalue()



def create_word(html_content):
    """Convert newsletter HTML into a downloadable Word document."""
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    soup = BeautifulSoup(html_content, "html.parser")

    for element in soup.find_all(
        ["h1", "h2", "h3", "h4", "p", "li", "blockquote"]
    ):
        text = element.get_text(" ", strip=True)
        if not text:
            continue

        if element.name == "h1":
            document.add_heading(text, level=1)
        elif element.name == "h2":
            document.add_heading(text, level=2)
        elif element.name in ("h3", "h4"):
            document.add_heading(text, level=3)
        elif element.name == "li":
            document.add_paragraph(text, style="List Bullet")
        elif element.name == "blockquote":
            document.add_paragraph(text, style="Quote")
        else:
            document.add_paragraph(text)

    # Include article links that may not be visible in paragraph text.
    links = soup.find_all("a", href=True)
    if links:
        document.add_heading("Article Links", level=2)
        seen = set()
        for link in links:
            url = link["href"].strip()
            label = link.get_text(" ", strip=True) or url
            if url.startswith(("http://", "https://")) and url not in seen:
                document.add_paragraph(f"{label}: {url}")
                seen.add(url)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
